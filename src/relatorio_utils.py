"""Funções auxiliares de formatação para gerar os relatórios em .docx."""
from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

import config

AZUL = RGBColor(0x1F, 0x3B, 0x73)
_MESES = ["jan", "fev", "mar", "abr", "mai", "jun",
          "jul", "ago", "set", "out", "nov", "dez"]


def mes_ano(data_iso: str) -> str:
    a, m, _ = data_iso.split("-")
    return f"{_MESES[int(m) - 1]}/{a}"


def data_br(data) -> str:
    return data.strftime("%d/%m/%Y")


def h1(doc, txt):
    p = doc.add_heading(txt, level=1)
    for r in p.runs:
        r.font.color.rgb = AZUL


def h2(doc, txt):
    p = doc.add_heading(txt, level=2)
    for r in p.runs:
        r.font.color.rgb = AZUL


def par(doc, txt, just=True):
    p = doc.add_paragraph(txt)
    if just:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def fig(doc, nome, legenda, largura=6.0):
    caminho = config.FIGURES_DIR / nome
    if caminho.exists():
        doc.add_picture(str(caminho), width=Inches(largura))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph(legenda)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cap.runs:
            r.font.size = Pt(9)
            r.font.italic = True


def tabela(doc, df, titulo, fmt=None):
    cap = doc.add_paragraph(titulo)
    for r in cap.runs:
        r.bold = True
        r.font.color.rgb = AZUL
    t = doc.add_table(rows=1, cols=len(df.columns) + 1)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    hdr[0].text = df.index.name or ""
    for j, c in enumerate(df.columns):
        hdr[j + 1].text = str(c)
    for idx, row in df.iterrows():
        cells = t.add_row().cells
        cells[0].text = str(idx)
        for j, c in enumerate(df.columns):
            v = row[c]
            if isinstance(v, float):
                cells[j + 1].text = (fmt(c, v) if fmt else f"{v:.3f}")
            else:
                cells[j + 1].text = str(v)
    doc.add_paragraph()
